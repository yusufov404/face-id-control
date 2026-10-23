"""Report building and export (Excel / PDF)."""
import io
from datetime import date, datetime
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_excel(
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence],
    *,
    meta: dict | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"

    r = 1
    ws.cell(row=r, column=1, value=title).font = Font(size=14, bold=True)
    r += 1
    if meta:
        for k, v in meta.items():
            ws.cell(row=r, column=1, value=f"{k}:").font = Font(bold=True)
            ws.cell(row=r, column=2, value=str(v))
            r += 1
    r += 1

    header_row = r
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    r += 1

    for row in rows:
        for c, value in enumerate(row, start=1):
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M")
            elif isinstance(value, date):
                value = value.isoformat()
            ws.cell(row=r, column=c, value=value)
        r += 1

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=header_row, column=c).column_letter].width = max(
            12, min(40, len(str(headers[c - 1])) + 6)
        )
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence],
    *,
    meta: dict | None = None,
    font_size: int = 8,
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=20, rightMargin=20, topMargin=24, bottomMargin=24,
        title=title,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"])]
    if meta:
        meta_text = " • ".join(f"{k}: {v}" for k, v in meta.items())
        story.append(Paragraph(meta_text, styles["Normal"]))
    story.append(Spacer(1, 10))

    data = [list(headers)]
    for row in rows:
        data.append([
            v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime)
            else v.isoformat() if isinstance(v, date)
            else "" if v is None else str(v)
            for v in row
        ])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B0BEC5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FB")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(table)
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        f"Yaratilgan: {datetime.now().strftime('%Y-%m-%d %H:%M')} • FaceSys Enterprise",
        styles["Italic"],
    ))
    doc.build(story)
    return buf.getvalue()